"""Import member names and phone numbers from the existing Excel layout."""

from collections import Counter
from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.contributions.members import normalize_sa_phone, register_member
from app.models.members import Member


MAX_IMPORT_BYTES = 3 * 1024 * 1024


class ImportError(ValueError):
    pass


def _header(value: object) -> str:
    return " ".join(str(value or "").upper().replace(":", "").split())


def _name_key(value: str) -> str:
    return " ".join(value.upper().split())


def import_members(db: Session, group_id: UUID, data: bytes) -> dict:
    if not data or len(data) > MAX_IMPORT_BYTES or not data.startswith(b"PK"):
        raise ImportError("Upload an XLSX workbook no larger than 3 MB")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError("Could not read the XLSX workbook") from exc

    rows: list[tuple[str, str | None]] = []
    for sheet in workbook.worksheets:
        header_row = number_col = name_col = phone_col = None
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20)):
            labels = [_header(cell.value) for cell in row]
            for index, label in enumerate(labels):
                if label in {"NO", "NUMBER"}:
                    number_col = index + 1
                if label == "MEMBERS":
                    header_row, name_col = row[0].row, index + 1
                if label in {"CONTACT NO", "CONTACT NUMBER", "PHONE NUMBER", "CELLPHONE NUMBER"}:
                    phone_col = index + 1
            if header_row and number_col and name_col and phone_col:
                break
        if not (header_row and number_col and name_col and phone_col):
            continue
        blank_run = 0
        for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            raw_number = values[number_col - 1] if len(values) >= number_col else None
            raw_name = values[name_col - 1] if len(values) >= name_col else None
            if not isinstance(raw_number, (int, float)) or raw_name is None or not str(raw_name).strip():
                blank_run += 1
                if blank_run >= 3:
                    break
                continue
            blank_run = 0
            name = " ".join(str(raw_name).split())
            raw_phone = values[phone_col - 1] if len(values) >= phone_col else None
            phone = None
            if raw_phone is not None and str(raw_phone).strip():
                try:
                    phone = normalize_sa_phone(str(raw_phone).split(".")[0])
                except ValueError as exc:
                    raise ImportError(f"Invalid contact number for member row {header_row + len(rows) + 1}") from exc
            rows.append((name, phone))
        if rows:
            break
    workbook.close()
    if not rows:
        raise ImportError("Could not find NO, MEMBERS, and CONTACT NO columns")

    existing = {
        _name_key(member.name)
        for member in db.scalars(select(Member).where(Member.group_id == group_id))
    }
    created = skipped = missing = 0
    imported_phones = []
    seen_names = set(existing)
    for name, phone in rows:
        key = _name_key(name)
        if key in seen_names:
            skipped += 1
            continue
        register_member(db, group_id, name, [], contact_number=phone)
        seen_names.add(key)
        created += 1
        if phone:
            imported_phones.append(phone)
        else:
            missing += 1
    counts = Counter(imported_phones)
    shared = sum(1 for count in counts.values() if count > 1)
    warnings = []
    if missing:
        warnings.append(f"{missing} imported members have no contact number.")
    if shared:
        warnings.append(f"{shared} contact number is shared by multiple imported members; use the POP reference to resolve it.")
    return {
        "created": created,
        "skipped_existing": skipped,
        "missing_contact_numbers": missing,
        "shared_contact_numbers": shared,
        "warnings": warnings,
    }
